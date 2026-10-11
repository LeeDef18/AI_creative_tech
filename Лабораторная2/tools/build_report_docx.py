"""Build the Word report from final Markdown artifacts (separate docx environment)."""
import argparse
import csv
import re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def inline(p, text):
    # Keep link destinations readable in the offline report.
    text = re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    for i, part in enumerate(re.split(r'(\*\*.*?\*\*|`[^`]+`)', text)):
        if not part:
            continue
        run = p.add_run(part.strip('*') if part.startswith('**') else part.strip('`'))
        if part.startswith('**'):
            run.bold = True
        elif part.startswith('`'):
            run.font.name = 'Consolas'
            run.font.size = Pt(10)


def table(doc, rows):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = 'Table Grid'
    t.autofit = False
    widths = {
        'Объект': [1.7, 2.1, 3.6, 9.6],
        'Проверка': [6.0, 1.8, 3.1, 6.1],
        'ID': [1.4, 11.1, 1.5, 1.5, 1.5],
        'Риск': [1.6, 1.1, 1.1, 1.1, 12.1],
        'Метрика': [2.5, 3.3, 2.6, 4.0, 4.6],
        'Поле политики': [6.0, 11.0],
    }.get(rows[0][0], [17 / len(rows[0])] * len(rows[0]))
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for idx, row in enumerate(rows):
        cells = t.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Cm(width)
            inline(cell.paragraphs[0], value)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.line_spacing = 1
                for r in p.runs:
                    r.font.size = Pt(10)
                    if idx == 0:
                        r.bold = True
        pr = t.rows[-1]._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if idx == 0:
            pr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def markdown(doc, text, skip_title=False):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                p = doc.add_paragraph(style='Code')
                p.add_run(lines[i])
                i += 1
        elif s.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', x.replace(' ', '')) for x in row):
                    rows.append(row)
                i += 1
            table(doc, rows)
            continue
        elif s.startswith('#'):
            m = re.match(r'(#+)\s+(.*)', s)
            if m and not (skip_title and len(m[1]) == 1):
                doc.add_heading(m[2], level=min(len(m[1]), 3))
        else:
            style = 'List Bullet' if s.startswith('- ') else None
            p = doc.add_paragraph(style=style)
            inline(p, s[2:] if style else s)
        i += 1


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--out', type=Path, default=ROOT/'reports/Отчёт_ЛР2_Вариант23_черновик.docx')
    args = parser.parse_args()
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.5)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Times New Roman', Pt(12)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(4)
    for name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[name].font.name = 'Times New Roman'
        doc.styles[name].font.color.rgb = __import__('docx').shared.RGBColor(0, 0, 0)
    code = doc.styles.add_style('Code', 1)
    code.font.name, code.font.size = 'Consolas', Pt(9)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1
    footer = sec.footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    sec.different_first_page_header_footer = True
    for text in ['Донской государственный технический университет', '(ДГТУ)']:
        p = doc.add_paragraph(text)
        p.alignment = 1
    doc.add_paragraph('\n\n')
    for text in ['ЛАБОРАТОРНАЯ РАБОТА № 2', 'Карточка модели, лицензии, происхождение данных и анализ рисков', 'Искусственный интеллект в креативных технологиях', 'Вариант 23']:
        p = doc.add_paragraph(text)
        p.alignment = 1
        p.runs[0].bold = True
    doc.add_paragraph('\n\n')
    for text in ['Выполнил: Король Алексей Юрьевич', 'Группа: Т26ИИКТ-МИМО11', 'Дата выполнения: 11.10.2026']:
        p = doc.add_paragraph(text)
        p.alignment = 2
    doc.add_paragraph('\n\n2026').alignment = 1
    doc.add_page_break()
    report = (ROOT/'reports/report.md').read_text()
    markdown(doc, report[report.index('## 1.'):])
    for title, filename in [('Приложение А. Паспорт набора данных', 'data_card.md'), ('Приложение Б. Карточка модели', 'model_card.md')]:
        doc.add_page_break()
        doc.add_heading(title, 1)
        markdown(doc, (ROOT/'artifacts/v23'/filename).read_text(), skip_title=True)
    doc.add_page_break()
    doc.add_heading('Приложение В. Реестр рисков', 1)
    with (ROOT/'artifacts/v23/risk_register.csv').open(newline='') as f:
        risks = list(csv.DictReader(f))
    for risk in risks:
        doc.add_heading(risk['risk_id'] + '. ' + risk['category'], 2)
        for label, key in [('Описание', 'description'), ('Мера', 'measure'), ('Ответственная роль', 'owner'), ('Признак наступления', 'trigger')]:
            p = doc.add_paragraph()
            p.add_run(label + ': ').bold = True
            p.add_run(risk[key])
        doc.add_paragraph(f"L = {risk['likelihood']}; I = {risk['impact']}; R = {int(risk['likelihood']) * int(risk['impact'])}.")
    markdown(doc, (ROOT/'reports/risk_rationale.md').read_text(), skip_title=True)
    doc.add_page_break()
    doc.add_heading('Приложение Г. Контрольные вопросы', 1)
    markdown(doc, (ROOT/'reports/control_questions.md').read_text(), skip_title=True)
    doc.core_properties.title = 'Лабораторная работа № 2. Вариант 23'
    doc.core_properties.author = 'Король Алексей Юрьевич'
    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    print(args.out)


if __name__ == '__main__':
    main()
