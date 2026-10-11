"""Build editable Word report from reports/report.md.
Requires python-docx==1.2.0 (report formatting only, outside experiment .venv).
"""
from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'reports/report.md'
OUT = ROOT / 'reports/Отчёт_ЛР1_Вариант23_черновик.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.5)
sec.different_first_page_header_footer = True
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Times New Roman', Pt(12)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.first_line_indent = Cm(0.75)
for name, size in [('Title',18),('Heading 1',14),('Heading 2',12)]:
    style = doc.styles[name]
    style.font.name, style.font.size = 'Times New Roman', Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title = 'Лабораторная работа № 1. Вариант 23'
doc.core_properties.author = ''
doc.core_properties.subject = 'Воспроизводимость и квантование текстуры «водная рябь»'


def inline(p, text):
    # Local references are written as human-readable paths for a standalone document.
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)',lambda m: m[1]+' ('+m[2]+')',text)
    for token in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)',text):
        if not token: continue
        if token.startswith('**'):
            run=p.add_run(token[2:-2]);run.bold=True
        elif token.startswith('`'):
            run=p.add_run(token[1:-1]);run.font.name='Consolas';run.font.size=Pt(10)
        else:p.add_run(token)


def para(text, style=None):
    p=doc.add_paragraph(style=style);inline(p,text)
    if style is None:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def caption(text):
    p=doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0)
    for run in p.runs:run.font.size=Pt(10);run.italic=True


def table(headers, rows, widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    if widths is None:widths=[17/len(headers)]*len(headers)
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for cell,w,text in zip(t.rows[0].cells,widths,headers):
        cell.width=Cm(w);cell.text=text
        shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'EAEAEA');cell._tc.get_or_add_tcPr().append(shading)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for vals in rows:
        cells=t.add_row().cells
        for cell,w,text in zip(cells,widths,vals):cell.width=Cm(w);cell.text=str(text)
    for ri,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0)
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                p.paragraph_format.line_spacing=1
                for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(10);run.bold=ri==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t


# Separate title page with editable placeholders.
for text in ['[Название учебного заведения]','[Институт / факультет]','[Кафедра]']:
    p=doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
doc.add_paragraph('\n\n')
p=doc.add_paragraph('ОТЧЁТ',style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for text in ['по лабораторной работе № 1','«Подготовка рабочей среды и паспорт воспроизводимого AI-эксперимента»','Дисциплина: «Искусственный интеллект в креативных технологиях»','Вариант 23 — водная рябь, квантование до 4 бит']:
    p=doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
doc.add_paragraph('\n')
for text in ['Выполнил: [ФИО]','Группа: [номер группы]','Проверил: [должность, ФИО преподавателя]']:
    p=doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.paragraph_format.first_line_indent=Cm(0)
p=doc.add_paragraph('\n\n[Город] — 2026');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent=Cm(0)
doc.add_page_break()

text=REPORT.read_text();text=text[text.index('## 1.'):]
lines=text.splitlines();i=0;table_no=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('## '):
        title=line[3:]
        if title.startswith('Приложение'):doc.add_page_break()
        doc.add_heading(title,level=1);i+=1;continue
    if line.startswith('### '):doc.add_heading(line[4:],level=2);i+=1;continue
    if line.startswith('```'):
        code=[];i+=1
        while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
        for c in code:
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(2)
            run=p.add_run(c);run.font.name='Consolas';run.font.size=Pt(9)
        i+=1;continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            block.append([x.strip().strip('`') for x in lines[i].strip().strip('|').split('|')]);i+=1
        headers=block[0];rows=block[2:]
        if any('![' in x for row in rows for x in row):
            t=doc.add_table(rows=1,cols=2);t.autofit=False
            for c,h,img in zip(t.rows[0].cells,headers,['base_seed101.png','perturbed_seed101.png']):
                c.width=Cm(8.5);p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
                p.add_run(h+'\n').bold=True;p.add_run().add_picture(str(ROOT/'artifacts/v23/artifacts'/img),width=Cm(6))
            caption('Рисунок 1 — Базовая и изменённая текстуры, seed=101. Оба исходных изображения 96×96 пикселей показаны в одинаковом масштабе.')
            continue
        table_no+=1;caption(f'Таблица {table_no}')
        if len(headers)==9:
            table([headers[j] for j in [0,1,2,3,4]],[[row[j] for j in [0,1,2,3,4]] for row in rows],[3,3.5,3,3.5,4])
            table([headers[j] for j in [0,5,6,7,8]],[[row[j] for j in [0,5,6,7,8]] for row in rows],[2.5,3,3,3,5.5])
        else:
            # Full SHA-256 values remain in the machine-readable artifacts.
            long_hash=any(len(x)==64 and re.fullmatch('[0-9a-f]+',x) for row in rows for x in row)
            if long_hash:
                rows=[[x[:16] if len(x)==64 and re.fullmatch('[0-9a-f]+',x) else x for x in row] for row in rows]
                headers=[x+' (первые 16 символов)' if 'SHA-256' in x else x for x in headers]
            table(headers,rows)
            if long_hash:para('Полные контрольные суммы сохранены в JSON-файлах результатов.')
        continue
    if line.startswith('- '):
        p=para(line[2:]);p.paragraph_format.first_line_indent=Cm(0);i+=1;continue
    paragraph=[line];i+=1
    while i<len(lines) and lines[i].strip() and not lines[i].startswith(('##','```','|','- ')):
        paragraph.append(lines[i].strip());i+=1
    para(' '.join(paragraph))

# Raw Julia output belongs in an appendix, leaving the main text readable.
doc.add_page_break();doc.add_heading('Приложение Б. Вывод среды Engee',level=1)
para('Вывод versioninfo() приведён без изменений. Результат отдельной команды VERSION: v"1.12.4".')
for line in (ROOT/'engee/versioninfo.txt').read_text().splitlines():
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
    run=p.add_run(line);run.font.name='Consolas';run.font.size=Pt(8)
doc.save(OUT)
print(OUT)
